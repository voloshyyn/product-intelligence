"""Word renderer. Findings are rendered with status and evidence, never bare facts."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
try:
    from .model import CATEGORIES, safe_asset
except ImportError:
    from model import CATEGORIES, safe_asset

LABELS={
 'title':('Product intelligence report','Исследование продукта'),
 'summary':('Executive summary','Резюме для принятия решений'),
 'method':('Methodology and limitations','Методология и ограничения'),
 'standard':('Product and market analysis','Анализ продукта и рынка'),
 'custom':('Custom research criteria','Дополнительные критерии исследования'),
 'ux':('Detailed UX audit','Подробный аудит UX'),
 'journeys':('Key user journeys','Ключевые пользовательские сценарии'),
 'hypotheses':('Hypotheses requiring validation','Гипотезы для проверки'),
 'opportunities':('Prioritized strategic opportunities','Приоритетные стратегические возможности'),
 'comparisons':('Comparable product evidence','Сравнение продуктов по доказательствам'),
 'conflicts':('Conflicting sources','Противоречия источников'),
 'sources':('Source registry and verification','Источники и сведения о проверке'),
 'screenshots':('Recorded screenshots','Зафиксированные снимки экрана'),
 'empty':('No findings recorded','Результаты не зафиксированы'),
 'company':('Company statement','Заявление компании'),
 'caveat':('Caveat','Оговорка'),
 'evidence':('Evidence','Доказательства'),
 'confidence':('Confidence','Уверенность'),
 'browser':('Browser testing','Браузерное тестирование'),
 'scope':('Scope','Область исследования'),
 'depth':('Depth','Глубина'),
 'run':('Run ID','Идентификатор запуска'),
 'date':('As of','По состоянию на'),
 'market':('Market','Рынок'),
 'criterion':('Criterion','Критерий'),
 'product':('Product','Продукт'),
 'claim':('Claim','Утверждение'),
 'status':('Status','Статус'),
 'evaluation':('Evaluation','Результат оценки'),
 'question':('Research question','Исследовательский вопрос'),
 'priority':('Priority','Приоритет'),
 'impact':('Expected impact','Ожидаемый эффект'),
 'effort':('Estimated effort','Оценка усилий'),
 'test':('Next validation step','Следующая проверка'),
 'severity':('Severity','Серьезность'),
 'goal':('User goal','Цель пользователя'),
 'surface':('Surface','Экран или раздел'),
 'heuristic':('Heuristic','Эвристика'),
 'frequency':('Frequency evidence','Свидетельства частоты'),
 'recommendation':('Recommendation','Рекомендация'),
 'action':('Action','Действие'),
 'expected':('Expected outcome','Ожидаемый результат'),
 'observed':('Observed outcome','Наблюдаемый результат'),
 'verification':('Verification status is distinct from confidence and provenance. Company statements establish what the company says. Missing evidence is a research gap.','Статус проверки отличается от уверенности и происхождения. Заявление компании подтверждает содержание ее сообщения. Отсутствие доказательств является пробелом исследования.'),
 'checklist':('Research checklist','План и полнота исследования'),
}


def labels_for(data):
    language=data['language'].split('-')[0]
    if language not in ('ru','en'):
        labels=data.get('report_labels',{})
        missing=(set(LABELS)|set(CATEGORIES))-set(labels)
        if missing:
            raise ValueError(f'Provide translated report_labels for {language}: {sorted(missing)}')
        return labels
    i=1 if language=='ru' else 0
    return {**{k:v[i] for k,v in LABELS.items()}, **{k:v[i] for k,v in CATEGORIES.items()}}


def link(paragraph,text,url):
    element=OxmlElement('w:hyperlink')
    element.set(qn('r:id'),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement('w:r')
    props=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'0563C1'); props.append(color)
    run.append(props)
    node=OxmlElement('w:t'); node.text=text; run.append(node)
    element.append(run); paragraph._p.append(element)


def table(doc,headers,records,widths=None):
    t=doc.add_table(rows=1,cols=len(headers))
    t.style='Table Grid'
    t.autofit=False
    for i,h in enumerate(headers):
        t.rows[0].cells[i].text=h
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for record in records:
        cells=t.add_row().cells
        for cell,value in zip(cells,record):
            cell.text=str(value)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+edge); node.set(qn('w:val'),'single'); node.set(qn('w:sz'),'4'); node.set(qn('w:color'),'D9D9D9'); borders.append(node)
    t._tbl.tblPr.append(borders)
    for n,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width=Inches(widths[i])
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'18354A' if n==0 else 'F3F6F8' if n%2==0 else 'FFFFFF'); props.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                m=OxmlElement('w:'+side); m.set(qn('w:w'),'100'); m.set(qn('w:type'),'dxa'); margins.append(m)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                for run in p.runs:
                    run.font.size=Pt(10)
                    if n==0:
                        run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)
    return t


def write_report(data,target,asset_root):
    labels=labels_for(data)
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.75)
    sec.left_margin=sec.right_margin=Inches(.8)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3'):
        style=doc.styles[name]
        style.font.name='Calibri'; style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Title'].font.size=Pt(30)
    doc.styles['Heading 1'].font.size=Pt(19)
    doc.styles['Heading 2'].font.size=Pt(14)
    doc.core_properties.title=labels['title']
    doc.core_properties.subject=data['scope']
    doc.core_properties.identifier=data['run_id']
    for style in ('Normal','Title','Heading 1','Heading 2'):
        lang=OxmlElement('w:lang'); lang.set(qn('w:val'),data['language']); doc.styles[style]._element.get_or_add_rPr().append(lang)
    footer=sec.footer.paragraphs[0]
    footer.add_run('Product intelligence  ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    doc.add_heading(labels['title'],0)
    doc.add_paragraph(' · '.join(p['name'] for p in data['products']),style='Subtitle')
    doc.add_paragraph(f'{labels["scope"]}: {data["scope"]}')
    for p in data['products']:
        paragraph=doc.add_paragraph(f'{p["name"]} · {p["industry"]} · {labels["market"]}: {p["market"]}\n')
        link(paragraph,p['url'],p['url'])
    doc.add_paragraph(f'{labels["date"]}: {data["completed_at"]}\n{labels["run"]}: {data["run_id"]}')
    doc.add_page_break()
    all_claims={c['claim_id']:c for c in data['claims']+data['hypotheses']}
    products={p['product_id']:p['name'] for p in data['products']}
    def finding(c):
        p=doc.add_paragraph()
        p.add_run(f'[{c["claim_id"]}] {c["status"]} · ').bold=True
        if c['company_statement']:
            p.add_run(labels['company']+': ').bold=True
        p.add_run(c['assertion'])
        if c.get('caveat'):
            doc.add_paragraph(f'{labels["caveat"]}: {c["caveat"]}')
        doc.add_paragraph(f'{labels["evidence"]}: {", ".join(c["source_ids"]+c["screenshot_ids"]) or "—"} · '
            f'{c["evidence_type"]} · {c["date_checked"]} · {labels["confidence"]}: {c["confidence"]:.0%}')
        if c['notes']:
            doc.add_paragraph(c['notes'])
    doc.add_heading(labels['summary'],1)
    for cid in data['summary_claim_ids']:
        c=all_claims[cid]
        if len(products)>1:
            doc.add_paragraph(products[c['product_id']],style='Heading 3')
        finding(c)
    if not data['summary_claim_ids']:
        doc.add_paragraph(labels['empty'])
    doc.add_heading(labels['method'],1)
    doc.add_paragraph(f'{labels["depth"]}: {data["depth"]}\n{labels["scope"]}: {data["scope"]}')
    doc.add_paragraph(data['plan']['source_strategy'])
    doc.add_paragraph(f'{labels["browser"]}: {data["tools"]["browser_tested"]} · {data["tools"]["browser_name"]}\n{data["tools"]["reason"]}')
    doc.add_paragraph(labels['verification'])
    for limitation in data['limitations']:
        doc.add_paragraph(limitation,style='List Bullet')
    doc.add_heading(labels['standard'],1)
    for p in data['products']:
        doc.add_heading(p['name'],2)
        for number,cat in enumerate(CATEGORIES,1):
            doc.add_heading(f'{number} {labels[cat]}',3)
            cov=next(r for r in data['coverage'] if r['product_id']==p['product_id'] and r['category']==cat)
            doc.add_paragraph(f'{cov["evaluation"]} · {cov["reason"]}')
            for c in data['claims']:
                if c['product_id']==p['product_id'] and c['category']==cat:
                    finding(c)
    doc.add_heading(labels['custom'],1)
    for criterion in data['custom_criteria']:
        doc.add_heading(criterion['label'],2)
        doc.add_paragraph(f'{labels["question"]}: {criterion["question"]}')
        for p in data['products']:
            r=next(r for r in data['custom_results'] if r['product_id']==p['product_id'] and r['criterion_id']==criterion['criterion_id'])
            doc.add_paragraph(f'{p["name"]}: {r["evaluation"]} · {r["reason"]}')
            for c in data['claims']:
                if c['product_id']==p['product_id'] and c.get('criterion_id')==criterion['criterion_id']:
                    finding(c)
    doc.add_heading(labels['ux'],1)
    for issue in sorted(data['ux_issues'],key=lambda r:-r['severity']):
        c=all_claims[issue['claim_id']]
        doc.add_heading(f'{issue["issue_id"]} {products[issue["product_id"]]}',2)
        finding(c)
        table(doc,[labels['severity'],labels['surface'],labels['heuristic']],[[f'{issue["severity"]}/4',issue['surface'],issue['heuristic']]],[.8,2.7,3.1])
        for key in ('user_goal','frequency_signal','recommendation'):
            label={'user_goal':'goal','frequency_signal':'frequency','recommendation':'recommendation'}[key]
            doc.add_paragraph(f'{labels[label]}: {issue[key]}')
    if not data['ux_issues']:
        doc.add_paragraph(labels['empty'])
    doc.add_heading(labels['journeys'],1)
    for j in data['journeys']:
        doc.add_heading(f'{j["journey_id"]} {products[j["product_id"]]}',2)
        doc.add_paragraph(f'{j["goal"]}\n{j["method"]} · {j["tool"]} · {j["tested_at"]} · {j["outcome"]}\n{j["limitations"]}')
        for step in sorted(j['steps'],key=lambda s:s['order']):
            doc.add_paragraph(f'{step["order"]} {labels["action"]}: {step["action"]}\n{labels["expected"]}: {step["expected"]}\n{labels["observed"]}: {step["observed"]}\n{labels["evidence"]}: {", ".join(step["claim_ids"])}')
    doc.add_heading(labels['hypotheses'],1)
    for c in data['hypotheses']:
        doc.add_paragraph(products[c['product_id']],style='Heading 3')
        finding(c)
    doc.add_heading(labels['opportunities'],1)
    for o in sorted(data['opportunities'],key=lambda r:(r['priority'],r['opportunity_id'])):
        doc.add_heading(f'{o["priority"]} {o["opportunity_id"]} {products[o["product_id"]]}',2)
        finding(all_claims[o['hypothesis_id']])
        for key,label in [('impact','impact'),('effort','effort'),('next_test','test')]:
            doc.add_paragraph(f'{labels[label]}: {o[key]}')
        doc.add_paragraph(f'{labels["evidence"]}: {", ".join(o["basis_claim_ids"])}')
    doc.add_heading(labels['comparisons'],1)
    for c in data['comparisons']:
        doc.add_heading(c['label'],2)
        # Compact metadata table; narrative stays in paragraphs to avoid narrow text columns.
        table(doc,[labels['product'],labels['claim'],labels['status']],[[products[x['product_id']],x['claim_id'],all_claims[x['claim_id']]['status']] for x in c['cells']],[2.7,1.4,2.5])
        for x in c['cells']:
            finding(all_claims[x['claim_id']])
        doc.add_paragraph(c['caveats'])
    doc.add_heading(labels['conflicts'],1)
    for c in data['conflicts']:
        doc.add_paragraph(f'{c["conflict_id"]}: {", ".join(c["claim_ids"])}\n{c["explanation"]}')
    doc.add_heading(labels['screenshots'],1)
    for shot in data['screenshots']:
        doc.add_heading(f'{shot["screenshot_id"]} {products[shot["product_id"]]}',2)
        picture=doc.add_picture(str(safe_asset(Path(asset_root),shot['path'])))
        ratio=min(Inches(6.5)/picture.width,Inches(7)/picture.height,1)
        picture.width=int(picture.width*ratio); picture.height=int(picture.height*ratio)
        doc.add_paragraph(f'{shot["caption"]}\n{shot["viewport"]} · {shot["captured_at"]}',style='Caption')
        link(doc.add_paragraph(),shot['source_url'],shot['source_url'])
    doc.add_heading(labels['sources'],1)
    for s in data['sources']:
        doc.add_heading(f'{s["source_id"]} {s["title"]}',2)
        link(doc.add_paragraph(),s['url'],s['url'])
        doc.add_paragraph(f'{products[s["product_id"]]} · {s["provenance"]} · {s["access"]} · {s["date_checked"]}\n{s["method"]} · {s["locator"]}')
        if s['excerpt']:
            doc.add_paragraph(s['excerpt'])
        if s['notes']:
            doc.add_paragraph(s['notes'])
    Path(target).parent.mkdir(parents=True,exist_ok=True)
    doc.save(target)
    # Read the saved OOXML package, not just the in-memory object.
    if Document(target).core_properties.identifier!=data['run_id']:
        raise ValueError('Word file verification failed')
