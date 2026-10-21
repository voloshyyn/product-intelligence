from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile
from docx import Document
from jsonschema import Draft202012Validator
from openpyxl import load_workbook, Workbook
from scripts.model import CATEGORIES, SHEETS, normalize_url, product_id, new_run, canonical_bytes, digest, slug
from scripts.schema_builder import build_schema
from scripts.validate_research import validate, load, ResearchError
from scripts.database import EXTRA_SHEETS, rows_for, run_path, storage_lock
from scripts.export_research import export, rebuild, verify_bundle
from tests.fixtures import fixture


class ResearchTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='product-intelligence-test-')
        self.addCleanup(self.tmp.cleanup)
        self.base=Path(self.tmp.name)
        self.assets=self.base/'assets'
        self.root=self.base/'isolated-storage'
        self.data=fixture(self.assets)

    def invalid(self,data=None):
        with self.assertRaises((ResearchError,ValueError)):
            validate(data or self.data,self.assets)

    def run_export(self,data=None):
        return export(data or self.data,self.root,self.assets)

    def test_schema_is_valid_and_matches_source(self):
        schema=load(Path(__file__).resolve().parents[1]/'references/research-schema.json')
        Draft202012Validator.check_schema(schema)
        self.assertEqual(schema,build_schema())
        validate(self.data,self.assets)

    def test_schema_missing_required_field(self):
        del self.data['claims'][0]['confidence']; self.invalid()

    def test_unknown_field_rejected(self):
        self.data['claims'][0]['revenue']=100; self.invalid()

    def test_invalid_status_rejected(self):
        self.data['claims'][0]['status']='probably'; self.invalid()

    def test_invalid_uuid_and_timezone(self):
        for field,value in [('run_id','../bad'),('started_at','2026-01-15T10:00:00')]:
            d=deepcopy(self.data); d[field]=value; self.invalid(d)

    def test_windows_case_collisions_rejected(self):
        d=deepcopy(self.data); d['run_id']='ABCDEF00-0000-4000-8000-000000000001'; self.invalid(d)
        duplicate=deepcopy(self.data['screenshots'][0])
        duplicate.update(screenshot_id='shot_uppercase',path='screenshots/FIXTURE.PNG')
        self.data['screenshots'].append(duplicate)
        with self.assertRaisesRegex(ResearchError,'Duplicate screenshot paths'): validate(self.data,self.assets)

    def test_missing_evidence_rejected(self):
        self.data['claims'][0]['source_ids']=[]; self.invalid()

    def test_dangling_source_rejected(self):
        self.data['claims'][0]['source_ids']=['missing']; self.invalid()

    def test_inaccessible_cannot_verify(self):
        self.data['sources'][0]['access']='NOT_ACCESSIBLE'; self.invalid()

    def test_official_statement_label_required(self):
        self.data['claims'][0]['company_statement']=False; self.invalid()

    def test_documentary_evidence_type_requires_correct_provenance(self):
        self.data['claims'][0]['evidence_type']='independent_report'; self.invalid()

    def test_partial_requires_caveat(self):
        c=next(c for c in self.data['claims'] if c['status']=='PARTIALLY_VERIFIED')
        del c['caveat']
        with self.assertRaisesRegex(ResearchError,'caveat'): validate(self.data,self.assets)

    def test_partial_requires_accessible_evidence(self):
        c=next(c for c in self.data['claims'] if c['status']=='PARTIALLY_VERIFIED')
        blocked=deepcopy(self.data['sources'][0]); blocked.update(source_id='denied',access='NOT_ACCESSIBLE')
        self.data['sources'].append(blocked); c['source_ids']=['denied']
        with self.assertRaisesRegex(ResearchError,'no accessible evidence'): validate(self.data,self.assets)

    def test_hypothesis_separation(self):
        h=self.data['hypotheses'].pop(); self.data['claims'].append(h); self.invalid()

    def test_unverified_never_has_factual_value(self):
        for status in ('UNVERIFIED','NOT_ACCESSIBLE','HYPOTHESIS'):
            d=deepcopy(self.data)
            c=d['hypotheses'][0] if status=='HYPOTHESIS' else d['claims'][2]
            c.update(status=status,value=999,facet='made_up'); self.invalid(d)

    def test_nonfinite_confidence(self):
        self.data['claims'][0]['confidence']=float('nan'); self.invalid()

    def test_source_dates_cannot_be_stale(self):
        self.data['sources'][0]['date_checked']='2020-01-01'; self.invalid()

    def test_future_completion(self):
        self.data['completed_at']='2099-01-01T00:00:00Z'; self.invalid()

    def test_screenshot_hash_and_missing_asset(self):
        self.data['screenshots'][0]['sha256']='0'*64; self.invalid()
        self.data['screenshots'][0]['path']='screenshots/missing.png'; self.invalid()

    def test_screenshot_traversal(self):
        self.data['screenshots'][0]['path']='screenshots/../escape.png'; self.invalid()

    def test_invalid_image_rejected_before_writing(self):
        image=self.assets/'screenshots/fixture.png'; image.write_bytes(b'not an image')
        self.data['screenshots'][0]['sha256']=hashlib.sha256(image.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ResearchError,'screenshot image'): self.run_export()
        self.assertFalse(self.root.exists())

    def test_duplicate_claim_id(self):
        self.data['claims'].append(deepcopy(self.data['claims'][0])); self.invalid()

    def test_cross_product_source(self):
        self.data['claims'][0]['source_ids']=['src_b']; self.invalid()

    def test_duplicate_product(self):
        self.data['products'].append(deepcopy(self.data['products'][0])); self.invalid()
        with self.assertRaises(ValueError): new_run(['https://x.example','https://x.example/?utm_source=z'])

    def test_url_identity_normalization_and_regions(self):
        self.assertEqual(normalize_url('HTTPS://EXAMPLE.COM:443/shop/?utm_source=a#hello'),'https://example.com/shop')
        self.assertEqual(product_id('https://example.com/?b=2&a=1'),product_id('https://example.com?a=1&b=2'))
        self.assertNotEqual(product_id('https://example.com','UA'),product_id('https://example.com','PL'))
        self.assertNotEqual(product_id('https://example.com/ua'),product_id('https://example.com/pl'))
        self.assertNotEqual(product_id('https://www.example.com'),product_id('https://example.com'))
        for url in ['javascript:alert(1)','https://a:b@example.com','https://example.com:bad']:
            with self.assertRaises(ValueError): normalize_url(url)

    def test_windows_safe_slugs(self):
        self.assertEqual(slug('CON'),'_con')
        self.assertNotIn(':',slug('retail:us/uk'))

    def test_all_categories_required(self):
        self.data['coverage'].pop(); self.invalid()

    def test_not_applicable_needs_reason_and_no_findings(self):
        self.data['coverage'][0].update(evaluation='NOT_APPLICABLE',reason='Physical product with no applicable interface')
        validate(self.data,self.assets)
        self.data['coverage'][0]['reason']=''; self.invalid()

    def test_custom_criteria_cannot_be_omitted(self):
        self.data['custom_results'].pop(); self.invalid()
        d=deepcopy(fixture(self.assets)); d['plan']['custom_criterion_ids']=[]; self.invalid(d)

    def test_arbitrary_custom_criterion(self):
        d=new_run(['https://example.com'],criteria=['Unusual criterion Ω arbitrary 🧭'])
        validate(d)
        self.assertEqual(d['custom_results'][0]['criterion_id'],'custom_1')

    def test_browser_unavailable_cannot_claim_testing(self):
        self.data['tools']['browser_available']=False; self.invalid()

    def test_browser_flag_requires_journey(self):
        self.data['journeys']=[]; self.invalid()

    def test_browser_steps_require_observations(self):
        self.data['journeys'][0]['steps'][0]['claim_ids']=['price_a']; self.invalid()

    def test_graceful_browser_degradation(self):
        d=new_run(['https://example.com'],criteria=['Shipping geography'])
        validate(d)
        result=self.run_export(d)
        self.assertTrue(Path(result['reports'][0]).is_file())

    def test_conflict_preservation(self):
        self.data['conflicts']=[{'conflict_id':'cf1','claim_ids':['price_a','price_b'],'explanation':'Conflicting synthetic versions'}]
        self.invalid()
        for c in self.data['claims']:
            if c['claim_id'] in ('price_a','price_b'):
                c.update(status='PARTIALLY_VERIFIED',caveat='Conflicting source; amount unresolved')
        validate(self.data,self.assets)
        self.assertEqual(len(rows_for([self.data],self.root)['Conflicts']),2)

    def test_summary_cannot_promote_hypothesis_or_unknown(self):
        for cid in ['hyp_a','unknown_a','does_not_exist']:
            d=deepcopy(self.data); d['summary_claim_ids']=[cid]; self.invalid(d)

    def test_comparison_category_and_identity(self):
        self.data['comparisons'][0]['cells'][1]['claim_id']='ux_a'; self.invalid()

    def test_word_and_screenshot_generation(self):
        r=self.run_export(); path=Path(r['reports'][0]); doc=Document(path)
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Резюме',text)
        for title in CATEGORIES.values(): self.assertIn(title[1],text)
        for criterion in self.data['custom_criteria']: self.assertIn(criterion['label'],text)
        self.assertIn('HYPOTHESIS',text)
        self.assertIn('Заявление компании',text)
        with ZipFile(path) as archive:
            self.assertTrue(any(n.startswith('word/media/') for n in archive.namelist()))
        verify_bundle(path.parent)

    def test_all_excel_sheets_and_formatting(self):
        r=self.run_export(); wb=load_workbook(r['workbook'])
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames,SHEETS+EXTRA_SHEETS)
        for ws in wb:
            self.assertEqual(ws.freeze_panes,'C2')
            self.assertTrue(ws.auto_filter.ref)
        ws=wb['Sources']; headers=[c.value for c in ws[1]]
        self.assertTrue(ws.cell(2,headers.index('url')+1).hyperlink)
        self.assertGreaterEqual(ws.row_dimensions[2].height,30)
        products=[r for r in rows_for([self.data],self.root)['Overview'] if r.get('record_type')=='product' and r['is_current']]
        self.assertEqual(len(products),2)

    def test_verified_fields_separate_company_and_partial_values(self):
        rows=rows_for([self.data],self.root)
        price=rows['Pricing & Business'][0]
        self.assertEqual(price['company_reported_value'],12)
        self.assertNotIn('verified_value',price)
        hypotheses=rows['Strengths & Weaknesses']
        self.assertEqual(hypotheses[0]['record_type'],'hypothesis')
        for records in rows.values():
            for row in records:
                if row.get('status') in ('UNVERIFIED','NOT_ACCESSIBLE','HYPOTHESIS'):
                    self.assertNotIn('verified_value',row)

    def test_verified_false_and_partial_zero_preserved(self):
        ux=next(c for c in self.data['claims'] if c['claim_id']=='ux_a')
        ux.update(facet='has_visible_label',value=False)
        partial=next(c for c in self.data['claims'] if c['status']=='PARTIALLY_VERIFIED')
        partial.update(facet='synthetic_limit',value=0)
        validate(self.data,self.assets)
        rows=rows_for([self.data],self.root)
        self.assertIs(rows['UX & UI'][0]['verified_value'],False)
        self.assertEqual(rows['Features & Services'][0]['partial_value'],0)

    def test_multi_industry_current_and_history(self):
        r=self.run_export()
        d=deepcopy(self.data); d['run_id']=str(uuid4())
        d['completed_at']='2026-01-15T11:00:00+00:00'
        d['claims'][0]['value']=18
        r2=self.run_export(d)
        self.assertTrue(Path(r['reports'][0]).is_file())
        self.assertNotEqual(r['reports'],r2['reports'])
        rows=rows_for([self.data,d],self.root)
        self.assertEqual({r['industry'] for r in rows['Research History']},{'Retail','Education'})
        self.assertEqual(len({r['run_id'] for r in rows['Research History']}),2)
        self.assertEqual(sum(r['is_current'] for r in rows['Pricing & Business']),2)
        self.assertEqual(len(list((self.root/'.snapshots').glob('*.json'))),2)

    def test_backdated_import_does_not_replace_current(self):
        newer=deepcopy(self.data); newer['run_id']=str(uuid4()); newer['completed_at']='2026-01-15T11:00:00Z'
        self.run_export(newer); self.run_export()
        current=[r for r in rows_for([newer,self.data],self.root)['Pricing & Business'] if r['is_current']]
        self.assertTrue(all(r['run_id']==newer['run_id'] for r in current))

    def test_idempotent_reexport_is_byte_preserving(self):
        r=self.run_export(); before={str(p):p.read_bytes() for p in self.root.rglob('*') if p.is_file()}
        r2=self.run_export()
        self.assertTrue(r2['idempotent'])
        self.assertEqual(before,{str(p):p.read_bytes() for p in self.root.rglob('*') if p.is_file()})

    def test_same_run_changed_content_rejected(self):
        self.run_export(); self.data['scope']='changed'
        with self.assertRaisesRegex(ResearchError,'different content'): self.run_export()

    def test_failure_preserves_workbook_and_rolls_back(self):
        self.run_export(); before=(self.root/'competitors.xlsx').read_bytes()
        d=deepcopy(self.data); d['run_id']=str(uuid4())
        from scripts.export_research import replace as actual_replace
        def fail(source,target):
            if Path(target).name=='competitors.xlsx': raise PermissionError('Synthetic Excel file lock')
            actual_replace(source,target)
        with patch('scripts.export_research.replace',side_effect=fail):
            with self.assertRaises(PermissionError): self.run_export(d)
        self.assertEqual(before,(self.root/'competitors.xlsx').read_bytes())
        self.assertEqual(len(list((self.root/'.snapshots').glob('*.json'))),1)
        self.assertFalse(run_path(self.root,d,d['products'][0]).exists())
        self.assertFalse((self.root/'.export.lock').exists())
        self.run_export(d)

    def test_bad_input_has_no_storage_side_effects(self):
        self.data['claims'][0]['source_ids']=[]
        with self.assertRaises(ResearchError): self.run_export()
        self.assertFalse(self.root.exists())

    def test_generation_error_preserves_existing_storage(self):
        self.run_export(); before=(self.root/'competitors.xlsx').read_bytes()
        d=deepcopy(self.data); d['run_id']=str(uuid4())
        with patch('scripts.export_research.write_report',side_effect=OSError('Synthetic disk error')):
            with self.assertRaises(OSError): self.run_export(d)
        self.assertEqual(before,(self.root/'competitors.xlsx').read_bytes())
        self.assertEqual(len(list((self.root/'.snapshots').glob('*.json'))),1)

    def test_unrelated_workbook_not_overwritten(self):
        self.root.mkdir(); Workbook().save(self.root/'competitors.xlsx')
        before=(self.root/'competitors.xlsx').read_bytes()
        with self.assertRaisesRegex(ResearchError,'unrelated'): self.run_export()
        self.assertEqual(before,(self.root/'competitors.xlsx').read_bytes())

    def test_corrupt_workbook_preserved(self):
        self.root.mkdir(); path=self.root/'competitors.xlsx'; path.write_bytes(b'damaged xlsx')
        with self.assertRaisesRegex(ResearchError,'corrupt'): self.run_export()
        self.assertEqual(path.read_bytes(),b'damaged xlsx')

    def test_stale_projection_requires_rebuild(self):
        result=self.run_export()
        wb=load_workbook(result['workbook']); wb.properties.keywords='old revision'; wb.save(result['workbook']); wb.close()
        with self.assertRaisesRegex(ResearchError,'stale'): self.run_export()
        rebuild(self.root)
        self.assertTrue(self.run_export()['idempotent'])

    def test_report_language_handling(self):
        from scripts.report import write_report, LABELS
        en=fixture(self.assets,multiple=False,language='en')
        path=self.base/'english.docx'; write_report(en,path,self.assets)
        self.assertIn('Executive summary',[p.text for p in Document(path).paragraphs])
        de=deepcopy(en); de['language']='de'
        with self.assertRaisesRegex(ValueError,'report_labels'): write_report(de,self.base/'german.docx',self.assets)
        de['report_labels']={k:'Übersetzung '+k for k in set(LABELS)|set(CATEGORIES)}
        write_report(de,self.base/'german.docx',self.assets)
        self.assertEqual(Document(self.base/'german.docx').paragraphs[0].text,'Übersetzung title')

    def test_current_identity_survives_industry_reclassification(self):
        self.run_export()
        d=deepcopy(self.data); d['run_id']=str(uuid4()); d['completed_at']='2026-01-15T11:00:00Z'
        d['products'][0]['industry']='Financial Services'
        self.run_export(d)
        current=[r for r in rows_for([self.data,d],self.root)['Overview'] if r.get('record_type')=='product' and r['is_current']]
        self.assertEqual(len(current),2)
        self.assertIn('Financial Services',{r['industry'] for r in current})

    def test_storage_lock(self):
        with storage_lock(self.root):
            with self.assertRaisesRegex(ResearchError,'locked'): self.run_export()
        self.assertFalse((self.root/'.export.lock').exists())

    def test_formula_injection_is_literal_text(self):
        self.data['claims'][0]['assertion']='=HYPERLINK("https://evil.example","click")'
        r=self.run_export(); wb=load_workbook(r['workbook']); self.addCleanup(wb.close)
        ws=wb['Pricing & Business']; col=[c.value for c in ws[1]].index('assertion')+1
        self.assertEqual(ws.cell(2,col).data_type,'s')
        self.assertEqual(ws.cell(2,col).value,self.data['claims'][0]['assertion'])

    def test_rebuild_and_tamper_detection(self):
        self.run_export(); (self.root/'competitors.xlsx').unlink()
        r=rebuild(self.root); self.assertEqual(r['runs'],1)
        folder=run_path(self.root,self.data,self.data['products'][0])
        (folder/'research.json').write_text('{}',encoding='utf-8')
        with self.assertRaisesRegex(ResearchError,'changed or incomplete'): rebuild(self.root)

    def test_duplicate_json_keys_rejected(self):
        path=self.base/'bad.json'; path.write_text('{"a":1,"a":2}',encoding='utf-8')
        with self.assertRaisesRegex(ResearchError,'Duplicate JSON key'): load(path)

    def test_cli_end_to_end_and_validation_failure(self):
        skill=Path(__file__).resolve().parents[1]
        plan=self.base/'plan.json'
        command=[sys.executable,str(skill/'scripts/init_research.py'),'--url','https://synthetic.example/',
                 '--criterion','Arbitrary support model','--output',str(plan)]
        result=subprocess.run(command,capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        result=subprocess.run([sys.executable,str(skill/'scripts/validate_research.py'),str(plan)],capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        result=subprocess.run([sys.executable,str(skill/'scripts/export_research.py'),str(plan),'--root',str(self.root)],capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        plan.write_text('{}',encoding='utf-8')
        result=subprocess.run([sys.executable,str(skill/'scripts/validate_research.py'),str(plan)],capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,2)


if __name__=='__main__':
    unittest.main()
